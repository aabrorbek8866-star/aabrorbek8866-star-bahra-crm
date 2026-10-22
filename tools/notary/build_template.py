# Notarius shablonini qayta yasash (public/templates/notarius_shablon.docx).
#
# Manba hujjatlar repozitoriyga QO'SHILMAGAN — ularda haqiqiy mijozlarning pasport ma'lumotlari bor.
# Shablonni o'zgartirish kerak bo'lsa: 4 ta Word faylni shu papkaga quyidagi nomlar bilan qo'ying,
# .doc bo'lsa avval .docx ga o'giring (LibreOffice), so'ng:
#   m_A.docx — Ma'lumotnoma (sof aktivlar)            m_B.docx — Ma'lumotnomalar (qarzdorlik yo'q, qoldiq qiymat)
#   m_C.docx — Umumiy yig'ilish bayonlari №1, №2      m_D.docx — Oldi-sotdi shartnomasi
#   pip install python-docx docxcompose
#   python build_template.py ../../public/templates/notarius_shablon.docx
# Qizil (FF0000) bilan yozilgan joylar {o'zgaruvchi}ga almashtiriladi; paragraf/run raqamlari aynan shu
# hujjatlarga mos — hujjat matni o'zgarsa, raqamlarni yangilash kerak (skript qolgan qizil joylarni chiqaradi).
# Mijozning qizil bilan yozilgan ma'lumotlarini {o'zgaruvchi}larga almashtirib, 4 hujjatni bitta shablonga birlashtiradi.
import sys, copy, docx
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_BREAK
from docxcompose.composer import Composer

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

def paras(d):
    return [Paragraph(p, d) for p in d.element.body.iter(W + 'p')]

def apply(d, edits, name):
    ps = paras(d)
    for (pi, ri), text in edits.items():
        runs = ps[pi].runs
        r = runs[ri]
        r.text = text
        r.font.color.rgb = None  # qizil → oddiy (qora) matn
    # nazorat: hech qanday qizil qolmasin (bo'sh runlardan tashqari)
    left = [(i, r.text) for i, p in enumerate(paras(d)) for r in p.runs
            if r.font.color is not None and r.font.color.type is not None and str(r.font.color.rgb) == "FF0000" and r.text.strip()]
    print(name, "qolgan qizil:", left)

def replace_in_run(d, pi, ri, old, new):
    r = paras(d)[pi].runs[ri]
    assert old in r.text, (pi, ri, r.text)
    r.text = r.text.replace(old, new, 1)

# ---------- D: oldi-sotdi shartnomasi ----------
D = docx.Document("m_D.docx")
body = D.element.body
kids = list(body)
# Birinchi shartnoma imzo jadvalidan keyingi hammasini (boshqa mijozlarning eski nusxalari) olib tashlash
end = next(k for k, el in enumerate(kids) if 'Рахматова' in ''.join(el.itertext()) and '___' in ''.join(el.itertext()))
for el in kids[end + 1:]:
    if el.tag != W + 'sectPr':
        body.remove(el)
apply(D, {
    (2, 1): "{shartnomaRaqami}",
    (4, 6): "{shartnomaSanasi}", (4, 7): "", (4, 8): "", (4, 9): "",
    (6, 5): "{fio}", (6, 8): "{pasportSeriya}", (6, 10): "{pasportRaqam}",
    (6, 16): "{podyezd}", (6, 18): "{qavat}", (6, 20): "{xonalar}", (6, 23): "{xonadon}",
    (6, 26): "{umumiyMaydon}", (6, 28): "{yashashMaydon} ",
    (9, 1): "{narxSoz}",
    (52, 0): "{fio}", (54, 1): "{pasportSeriya}", (54, 3): "№{pasportRaqam}",
    (56, 0): "{pasportKimTomonidan}",
    (60, 3): "{manzil}",
    (69, 6): "{imzo}", (69, 7): "", (69, 8): "", (69, 9): "",
}, "D")
# Pasport kim tomonidan va qachon berilgani (qora matn, lekin har mijozda boshqa)
replace_in_run(D, 6, 11, " Бухоро  вилояти Гиждивон тумани     ИИБ томонидан 23.03.2014 йилда",
               " {pasportKimTomonidan} томонидан {pasportSanasi} йилда")

# ---------- C: umumiy yig'ilish bayonlari ----------
C = docx.Document("m_C.docx")
apply(C, {
    (3, 2): "«{kun}» ", (3, 3): "{oy}", (3, 4): "",
    (31, 1): "«{kun}", (31, 2): "", (31, 3): "» {oy}",
    (37, 3): "{xonadon}",
    (39, 7): " {xonadon}", (39, 9): "{narxSoz}", (39, 12): "{fioGa}",
    (43, 3): "{xonadon}", (43, 5): "{narxSoz} ", (43, 8): "{fioGa}",
}, "C")
replace_in_run(C, 3, 5, "2026", "{yil}")
replace_in_run(C, 31, 4, "2026", "{yil}")

# ---------- B: ma'lumotnomalar (qarzdorlik yo'q, qoldiq qiymat) ----------
B = docx.Document("m_B.docx")
apply(B, {
    (3, 1): "№{xonadon}", (3, 2): "-", (3, 5): "{fioGa}", (3, 13): "{kadastr}",
    (3, 15): "{podyezd}", (3, 17): "{qavat}", (3, 19): "{xonalar}", (3, 21): "{xonadon}",
    (3, 23): "{shartnomaSanasi}", (3, 26): "{shartnomaRaqami}",
    (4, 1): "{kun}", (4, 2): "", (4, 3): "»", (4, 4): " ", (4, 5): "{oy}", (4, 6): "",
    (21, 4): "«{kun}» {oy}", (21, 13): "{xonadon}", (21, 18): "{narxSoz}",
}, "B")

replace_in_run(B, 4, 8, "202", "{yil}"); replace_in_run(B, 4, 9, "6", "")
replace_in_run(B, 21, 5, "2026", "{yil}")

# ---------- A: ma'lumotnoma (sof aktivlar) ----------
A = docx.Document("m_A.docx")
apply(A, {(14, 8): "«{kun}»", (14, 9): " ", (14, 10): "{oy} "}, "A")
replace_in_run(A, 14, 11, "202", "{yil}"); replace_in_run(A, 14, 12, "6", "")

for name, d in (("D", D), ("C", C), ("B", B), ("A", A)):
    d.save(f"t_{name}.docx")

# ---------- Birlashtirish: Shartnoma → Bayonlar → Ma'lumotnomalar ----------
from docx.shared import Pt

def effective_size(par, doc):
    st = par.style
    while st is not None:
        if st.font.size: return st.font.size
        st = st.base_style
    sz = doc.styles.element.find(f"{W}docDefaults/{W}rPrDefault/{W}rPr/{W}sz")
    return Pt(int(sz.get(f"{W}val")) / 2) if sz is not None else None

def effective_font(par, doc):
    st = par.style
    while st is not None:
        if st.font.name: return st.font.name
        st = st.base_style
    rf = doc.styles.element.find(f"{W}docDefaults/{W}rPrDefault/{W}rPr/{W}rFonts")
    return rf.get(f"{W}ascii") if rf is not None else None

def freeze_fonts(doc):
    """Qo'shilayotgan hujjatning shrift o'lchami va nomini runlarga yozib qo'yish — birlashtirilganda asosiy hujjat uslubi ustidan yozilmasin."""
    for par in paras(doc):
        size, font = effective_size(par, doc), effective_font(par, doc)
        runs = par.runs
        for i, r in enumerate(runs):
            if r.font.size is None:
                # Qo'shni matn o'lchami (o'zgaruvchi qo'yilgan bo'lak atrofidagi matn bilan bir xil bo'lsin)
                near = next((x.font.size for x in runs[i - 1::-1] if x.font.size), None) if i else None
                near = near or next((x.font.size for x in runs[i + 1:] if x.font.size), None)
                if near or size: r.font.size = near or size
            if r.font.name is None and font: r.font.name = font

def page_break_before(doc, pi):
    paras(doc)[pi].paragraph_format.page_break_before = True

master = docx.Document("t_D.docx")
comp = Composer(master)
for name, breaks in (("C", [28]), ("B", [19]), ("A", [])):
    nxt = docx.Document(f"t_{name}.docx")
    freeze_fonts(nxt)
    for pi in breaks:
        page_break_before(nxt, pi)
    comp.doc.add_page_break()  # har bir hujjat yangi betdan
    comp.append(nxt)
comp.save(sys.argv[1])
print("saqlandi:", sys.argv[1])
